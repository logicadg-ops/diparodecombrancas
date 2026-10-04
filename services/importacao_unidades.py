from io import BytesIO

from openpyxl import Workbook, load_workbook

from services import supabase_client

CAMPOS_TEXTO = [
    "bloco",
    "apto",
    "unidade",
    "proprietario",
    "cpfcnpj",
    "email",
    "fone1",
    "fone2",
    "rg",
    "inquilono",
    "inquilonocpfcnpj",
    "inquilinoemail",
    "inquilinofone1",
    "inquilinofone2",
]
CAMPO_BOOLEANO = "alugado"
CAMPOS = CAMPOS_TEXTO + [CAMPO_BOOLEANO]

_VERDADEIROS = {"sim", "s", "true", "1", "yes", "y", "x"}
_FALSOS = {"nao", "não", "n", "false", "0", "no"}


def gerar_modelo():
    wb = Workbook()
    ws = wb.active
    ws.title = "unidades"
    ws.append(CAMPOS)
    ws.append(
        ["01", "A", "01-LT-001", "Nome do proprietário", "12345678901", "email@exemplo.com",
         "81999990000", "", "", "não", "", "", "", "", ""]
    )
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _normalizar_texto(valor):
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    return texto or None


def _normalizar_booleano(valor):
    if valor is None or str(valor).strip() == "":
        return None
    texto = str(valor).strip().lower()
    if texto in _VERDADEIROS:
        return True
    if texto in _FALSOS:
        return False
    raise ValueError(f"valor inválido para alugado: '{valor}' (use sim ou não)")


def ler_planilha(arquivo_stream):
    """Lê a planilha e devolve (linhas, erros, avisos). Cada linha é um dict
    {numero_linha, dados} só com os campos preenchidos. Erros bloqueiam a
    importação; avisos (ex: colunas desconhecidas) não."""
    wb = load_workbook(arquivo_stream, read_only=True, data_only=True)
    ws = wb.active
    linhas_brutas = list(ws.iter_rows(values_only=True))
    wb.close()

    if not linhas_brutas:
        return [], ["A planilha está vazia."], []

    cabecalho = [str(c).strip().lower() if c is not None else "" for c in linhas_brutas[0]]
    if "unidade" not in cabecalho:
        return [], ["A planilha precisa ter uma coluna 'unidade' na primeira linha."], []

    colunas_desconhecidas = [c for c in cabecalho if c and c not in CAMPOS]
    avisos = []
    if colunas_desconhecidas:
        avisos.append(
            "Colunas ignoradas (não existem no cadastro): " + ", ".join(colunas_desconhecidas)
        )
    erros = []

    linhas = []
    for numero, linha in enumerate(linhas_brutas[1:], start=2):
        if not any(v is not None and str(v).strip() for v in linha):
            continue
        dados = {}
        try:
            for indice, nome_coluna in enumerate(cabecalho):
                if nome_coluna not in CAMPOS or indice >= len(linha):
                    continue
                valor = linha[indice]
                if nome_coluna == CAMPO_BOOLEANO:
                    valor_norm = _normalizar_booleano(valor)
                else:
                    valor_norm = _normalizar_texto(valor)
                if valor_norm is not None:
                    dados[nome_coluna] = valor_norm
        except ValueError as e:
            erros.append(f"Linha {numero}: {e}")
            continue
        linhas.append({"numero_linha": numero, "dados": dados})

    return linhas, erros, avisos


def planejar(id_condominio, nome_condominio, linhas, unidades_existentes):
    """Decide, linha a linha, se cria ou atualiza. Campos em branco na planilha
    não apagam dados já cadastrados (só campos preenchidos são aplicados)."""
    existentes_por_chave = {}
    for u in unidades_existentes:
        chave = (u.get("unidade") or "").strip()
        existentes_por_chave.setdefault(chave, []).append(u)

    vistos_no_arquivo = set()
    criar, atualizar, sem_alteracao, erros = [], [], [], []

    for linha in linhas:
        dados = linha["dados"]
        numero = linha["numero_linha"]
        chave = dados.get("unidade")

        if not chave:
            erros.append(f"Linha {numero}: coluna 'unidade' vazia.")
            continue
        if chave in vistos_no_arquivo:
            erros.append(f"Linha {numero}: unidade '{chave}' repetida na planilha.")
            continue
        vistos_no_arquivo.add(chave)

        if chave in existentes_por_chave:
            ids = [u["id"] for u in existentes_por_chave[chave]]
            alteracoes = {}
            for campo, novo in dados.items():
                antigos = {str(u.get(campo)) for u in existentes_por_chave[chave]}
                if str(novo) not in antigos:
                    alteracoes[campo] = novo
            if not alteracoes:
                sem_alteracao.append(chave)
            else:
                atualizar.append({"unidade": chave, "ids": ids, "dados": alteracoes})
        else:
            criar.append({
                **dados,
                "id_condominio": id_condominio,
                "nome_condominio": nome_condominio,
            })

    return {
        "criar": criar,
        "atualizar": atualizar,
        "sem_alteracao": sem_alteracao,
        "erros": erros,
    }


def aplicar(plano):
    if plano["criar"]:
        supabase_client.criar_unidades(plano["criar"])
    for item in plano["atualizar"]:
        for id_unidade in item["ids"]:
            supabase_client.atualizar_unidade_por_id(id_unidade, item["dados"])
